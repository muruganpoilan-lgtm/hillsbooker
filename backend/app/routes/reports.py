"""Reports: totals for a date range, and the same data as an Excel file. Admin sees every host; owners and operators see their own."""
import io
from datetime import date, datetime, timedelta
from decimal import Decimal
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response

from ..auth import require_role
from ..db import get_conn

router = APIRouter(prefix="/reports", tags=["reports"])
host_or_admin = require_role("owner", "operator", "admin")
IST = ZoneInfo("Asia/Kolkata")
ONLINE = ("razorpay", "cashfree")

_SQL = """
SELECT pm.id, (pm.created_at AT TIME ZONE 'Asia/Kolkata')::date AS paid_on, pm.provider, pm.status AS pay_status, pm.amount,
       pm.refunded_amount, pm.commission, pm.payout_id, po.paid_at AS payout_paid_at, po.reference AS payout_reference,
       s.kind, s.item, s.host_id, h.name AS host_name, h.phone AS host_phone, h.upi_id AS host_upi, s.guest, s.guest_phone, s.start_date, s.end_date, s.total_amount, s.status AS booking_status
FROM payments pm
JOIN (
  SELECT 'Stay' AS kind, b.id, b.status, p.name AS item, g.name AS guest, g.phone AS guest_phone, b.check_in AS start_date, b.check_out AS end_date,
         b.total_amount, p.owner_id AS host_id
  FROM bookings b JOIN properties p ON p.id = b.property_id JOIN users g ON g.id = b.user_id
  UNION ALL
  SELECT 'Vehicle', b.id, b.status, v.name, g.name, g.phone, b.start_date, b.end_date, b.total_amount, v.operator_id
  FROM vehicle_bookings b JOIN vehicles v ON v.id = b.vehicle_id JOIN users g ON g.id = b.user_id
) s ON (s.kind = 'Stay' AND pm.booking_id = s.id) OR (s.kind = 'Vehicle' AND pm.vehicle_booking_id = s.id)
JOIN users h ON h.id = s.host_id
LEFT JOIN payouts po ON po.id = pm.payout_id
WHERE pm.status IN ('paid', 'refunded')
  AND (pm.created_at AT TIME ZONE 'Asia/Kolkata')::date BETWEEN %s AND %s
  AND (%s OR s.host_id = %s)
ORDER BY pm.created_at
"""


def _period(date_from: str | None, date_to: str | None) -> tuple[date, date]:
    today = datetime.now(IST).date()
    try:
        start = date.fromisoformat(date_from) if date_from else today.replace(day=1)
        end = date.fromisoformat(date_to) if date_to else today
    except ValueError:
        raise HTTPException(400, "Dates must look like 2026-10-30")
    if end < start:
        raise HTTPException(400, "The end date is before the start date")
    if end - start > timedelta(days=366):
        raise HTTPException(400, "Pick a range of one year or less")
    return start, end


def _load(user: dict, start: date, end: date) -> list[dict]:
    with get_conn() as conn, conn.cursor() as cur:
        cur.execute(_SQL, (start, end, user["role"] == "admin", user["id"]))
        return cur.fetchall()


def _share(r: dict) -> Decimal:
    """The host's part of an online payment: what the guest paid, minus Hill Booker's fee, minus any refund."""
    return Decimal(r["amount"]) - Decimal(r["commission"]) - Decimal(r["refunded_amount"] or 0)


def _totals(rows: list[dict]) -> dict:
    paid = [r for r in rows if r["pay_status"] == "paid"]
    online = [r for r in paid if r["provider"] in ONLINE]
    zero = Decimal(0)
    return {
        "payments": len(paid),
        "advances_received": sum((Decimal(r["amount"]) for r in paid), zero),
        "refunded": sum((Decimal(r["refunded_amount"] or 0) for r in rows), zero),
        "commission": sum((Decimal(r["commission"]) for r in paid), zero),
        "fee_owed_by_hosts": sum((Decimal(r["commission"]) for r in paid if r["provider"] == "upi_direct"), zero),
        "due_to_hosts": sum((_share(r) for r in online if not r["payout_id"]), zero),
        "paid_out": sum((_share(r) for r in online if r["payout_id"]), zero),
    }


def _by_host(rows: list[dict]) -> list[dict]:
    groups: dict[str, list[dict]] = {}
    for r in rows:
        groups.setdefault(str(r["host_id"]), []).append(r)
    out = []
    for rs in groups.values():
        t = _totals(rs)
        t["host_name"] = rs[0]["host_name"]
        out.append(t)
    return sorted(out, key=lambda t: t["host_name"])


@router.get("/summary")
def summary(date_from: str | None = Query(None, alias="from"), date_to: str | None = Query(None, alias="to"), user: dict = Depends(host_or_admin)):
    start, end = _period(date_from, date_to)
    rows = _load(user, start, end)
    return {"from": start, "to": end, "totals": _totals(rows), "by_host": _by_host(rows) if user["role"] == "admin" else []}


def _payout_text(r: dict) -> str:
    if r["pay_status"] != "paid":
        return "Refunded"
    if r["provider"] not in ONLINE:
        return f"Guest paid the host directly. Host owes Hill Booker Rs {Decimal(r['commission']):,.2f}"
    if r["payout_id"]:
        return f"Paid to host on {r['payout_paid_at'].astimezone(IST):%d %b %Y}"
    return "Due after check-in" if r["start_date"] > datetime.now(IST).date() else "Due now"


@router.get("/export.xlsx")
def export_xlsx(date_from: str | None = Query(None, alias="from"), date_to: str | None = Query(None, alias="to"), user: dict = Depends(host_or_admin)):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    start, end = _period(date_from, date_to)
    rows = _load(user, start, end)
    admin = user["role"] == "admin"
    wb = Workbook()
    font, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    head_fill = PatternFill("solid", fgColor="1F4E79")
    head_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    thin = Side(style="thin", color="CCCCCC")
    money_fmt = '#,##0.00;(#,##0.00);"-"'

    def table(ws, headers, data, money_cols=(), date_cols=(), widths=None):
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill, c.alignment = head_font, head_fill, Alignment(vertical="center", wrap_text=True)
        for row in data:
            ws.append(row)
        for r in ws.iter_rows(min_row=2):
            for c in r:
                c.font = font
                c.border = Border(bottom=thin)
                if c.column in money_cols:
                    c.number_format = money_fmt
                if c.column in date_cols:
                    c.number_format = "DD-MMM-YYYY"
        for i, h in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, max(12, min(34, len(h) + 4)))
        ws.freeze_panes = "A2"
        if data:
            ws.auto_filter.ref = ws.dimensions

    # Summary sheet
    t = _totals(rows)
    ws = wb.active
    ws.title = "Summary"
    ws["A1"], ws["A1"].font = "Hill Booker report", Font(name="Arial", size=14, bold=True)
    ws["A2"], ws["A2"].font = f"Payments dated {start:%d %b %Y} to {end:%d %b %Y} (India time)", font
    ws["A3"], ws["A3"].font = ("All hosts" if admin else "Your bookings only"), font
    lines = [
        ("Paid bookings", t["payments"], False),
        ("Advances received", t["advances_received"], True),
        ("Refunded to guests", t["refunded"], True),
        ("Hill Booker commission earned", t["commission"], True),
        ("Due to hosts (paid online, not yet paid out)", t["due_to_hosts"], True),
        ("Paid out to hosts", t["paid_out"], True),
        ("Commission hosts owe (guest paid host by UPI)", t["fee_owed_by_hosts"], True),
    ]
    for i, (label, val, is_money) in enumerate(lines, start=5):
        ws.cell(i, 1, label).font = bold
        c = ws.cell(i, 2, float(val) if is_money else val)
        c.font = font
        if is_money:
            c.number_format = money_fmt
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 18
    ws.cell(13, 1, "Amounts in rupees. Dates are payment dates. Host share = advance minus commission minus refunds.").font = Font(name="Arial", size=9, italic=True)

    # Payments sheet
    heads = ["Date", "Type", "Property or vehicle", "Host", "Host phone", "Host UPI ID", "Guest", "Guest phone", "Check-in", "Check-out", "Booking total", "Advance paid", "How paid",
             "Payment status", "Refunded", "Commission", "Host share", "Payout", "Payout reference"]
    data = []
    for r in rows:
        online = r["provider"] in ONLINE
        data.append([
            r["paid_on"], r["kind"], r["item"], r["host_name"], r["host_phone"] or "", r["host_upi"] or "", r["guest"], r["guest_phone"] or "", r["start_date"], r["end_date"],
            float(r["total_amount"]), float(r["amount"]), "Paid online" if online else "UPI to host", r["pay_status"].capitalize(),
            float(r["refunded_amount"] or 0), float(r["commission"]), float(_share(r)) if online and r["pay_status"] == "paid" else 0.0,
            _payout_text(r), r["payout_reference"] or "",
        ])
    table(wb.create_sheet("Payments"), heads, data, money_cols=(11, 12, 15, 16, 17), date_cols=(1, 9, 10), widths={3: 30, 4: 22, 5: 16, 6: 22, 7: 22, 8: 16, 18: 44, 19: 18})

    if admin:
        bh = _by_host(rows)
        table(
            wb.create_sheet("By host"),
            ["Host", "Paid bookings", "Advances received", "Refunded", "Commission", "Due to host", "Paid out", "Host owes commission"],
            [[b["host_name"], b["payments"], float(b["advances_received"]), float(b["refunded"]), float(b["commission"]), float(b["due_to_hosts"]),
              float(b["paid_out"]), float(b["fee_owed_by_hosts"])] for b in bh],
            money_cols=(3, 4, 5, 6, 7, 8), widths={1: 28},
        )

    buf = io.BytesIO()
    wb.save(buf)
    name = f"hill-booker-report-{start:%Y%m%d}-{end:%Y%m%d}.xlsx"
    return Response(
        buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )
